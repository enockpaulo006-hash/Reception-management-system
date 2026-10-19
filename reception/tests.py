from django.contrib.auth.models import Group, User
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.management import call_command
from django.db import IntegrityError, connection
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from datetime import timedelta

from .models import AuditLog, DayClose, Payment, Service, Setting, Visit
from .services import next_receipt_no, record_payment, register_reprint, void_payment


class PaymentServiceTests(TestCase):
    def setUp(self):
        self.receptionist = User.objects.create_user("receptionist")
        self.owner = User.objects.create_user("owner")
        self.owner.groups.add(Group.objects.create(name="Owner"))
        self.visit = Visit.objects.create(
            name="Asha", age=30, sex="F", created_by=self.receptionist
        )
        self.consultation = Service.objects.create(
            name="Consultation", kind="CONSULT", price=10000
        )
        self.test_service = Service.objects.create(
            name="Blood test", kind="TEST", price=5000
        )

    def make_payment(self):
        return record_payment(
            self.visit, [self.consultation, self.test_service], self.receptionist
        )

    def test_receipt_numbers_are_sequential(self):
        self.assertEqual(next_receipt_no(), 1)
        self.assertEqual(next_receipt_no(), 2)

    def test_total_is_calculated_from_price_list(self):
        payment = self.make_payment()
        self.assertEqual(payment.total, 15000)
        self.assertEqual(list(payment.items.values_list("price", flat=True)), [10000, 5000])

    def test_payment_item_prices_are_frozen_after_price_change(self):
        payment = self.make_payment()
        self.consultation.price = 12500
        self.consultation.save()
        item = payment.items.get(service=self.consultation)
        self.assertEqual(item.name, "Consultation")
        self.assertEqual(item.price, 10000)

    def test_empty_payment_is_rejected(self):
        with self.assertRaises(ValidationError):
            record_payment(self.visit, [], self.receptionist)

    def test_receptionist_cannot_void(self):
        with self.assertRaises(PermissionDenied):
            void_payment(self.make_payment(), self.receptionist, "Wrong selection")

    def test_void_requires_reason_and_is_logged(self):
        payment = self.make_payment()
        with self.assertRaises(ValidationError):
            void_payment(payment, self.owner, "   ")
        void_payment(payment, self.owner, "Wrong selection")
        payment.refresh_from_db()
        self.assertTrue(payment.voided)
        self.assertEqual(payment.void_reason, "Wrong selection")
        self.assertTrue(AuditLog.objects.filter(action="VOID", user=self.owner).exists())

    def test_model_delete_is_blocked(self):
        with self.assertRaises(PermissionDenied):
            self.make_payment().delete()

    def test_model_amount_edit_is_blocked(self):
        payment = self.make_payment()
        payment.total = 1
        with self.assertRaises(PermissionDenied):
            payment.save()

    def test_raw_sql_delete_and_update_are_blocked_by_triggers(self):
        payment = self.make_payment()
        item = payment.items.first()
        audit = AuditLog.objects.filter(action="PAY").first()
        with connection.cursor() as cursor:
            with self.assertRaises(IntegrityError):
                cursor.execute("DELETE FROM reception_payment WHERE id = %s", [payment.pk])
            with self.assertRaises(IntegrityError):
                cursor.execute("DELETE FROM reception_paymentitem WHERE id = %s", [item.pk])
            with self.assertRaises(IntegrityError):
                cursor.execute("DELETE FROM reception_auditlog WHERE id = %s", [audit.pk])
            with self.assertRaises(IntegrityError):
                cursor.execute("UPDATE reception_payment SET total = 1 WHERE id = %s", [payment.pk])
            with self.assertRaises(IntegrityError):
                cursor.execute("UPDATE reception_paymentitem SET price = 1 WHERE id = %s", [item.pk])
            with self.assertRaises(IntegrityError):
                cursor.execute("UPDATE reception_auditlog SET action = 'EDIT' WHERE id = %s", [audit.pk])

    def test_reprint_is_counted_and_logged(self):
        payment = self.make_payment()
        register_reprint(payment, self.receptionist)
        payment.refresh_from_db()
        self.assertEqual(payment.reprints, 1)
        self.assertTrue(AuditLog.objects.filter(action="REPRINT", user=self.receptionist).exists())

    def test_dashboard_requires_login(self):
        response = self.client.get(reverse("dashboard"))
        self.assertRedirects(response, f"{reverse('login')}?next={reverse('dashboard')}")

    def test_dashboard_shows_todays_totals_and_voided_receipt(self):
        payment = self.make_payment()
        void_payment(payment, self.owner, "Entry made in error")
        self.client.force_login(self.receptionist)

        response = self.client.get(reverse("dashboard"))

        self.assertContains(response, "Today’s dashboard")
        self.assertContains(response, "TZS 0")
        self.assertContains(response, "VOIDED")
        self.assertContains(response, self.visit.name)

    def test_logout_only_accepts_post(self):
        self.client.force_login(self.receptionist)
        self.assertEqual(self.client.get(reverse("logout")).status_code, 405)
        response = self.client.post(reverse("logout"))
        self.assertRedirects(response, reverse("login"))

    def test_receptionist_cannot_void_through_view(self):
        payment = self.make_payment()
        self.client.force_login(self.receptionist)

        response = self.client.post(
            reverse("receipt_void", args=[payment.receipt_no]),
            {"reason": "Wrong service"},
            follow=True,
        )

        payment.refresh_from_db()
        self.assertFalse(payment.voided)
        self.assertContains(response, "Only the owner can void a receipt")

    def test_owner_can_void_through_view(self):
        payment = self.make_payment()
        self.client.force_login(self.owner)

        response = self.client.post(
            reverse("receipt_void", args=[payment.receipt_no]),
            {"reason": "Wrong service"},
        )

        self.assertRedirects(response, reverse("receipt_detail", args=[payment.receipt_no]))
        payment.refresh_from_db()
        self.assertTrue(payment.voided)
        self.assertEqual(payment.voided_by, self.owner)

    def test_reprint_view_registers_copy_and_redirects_for_printing(self):
        payment = self.make_payment()
        self.client.force_login(self.receptionist)

        response = self.client.post(reverse("receipt_reprint", args=[payment.receipt_no]))

        self.assertRedirects(
            response,
            f"{reverse('receipt_detail', args=[payment.receipt_no])}?copy=1&print=1",
        )
        payment.refresh_from_db()
        self.assertEqual(payment.reprints, 1)

    def test_day_close_hides_expected_cash_from_receptionist_before_submission(self):
        self.make_payment()
        self.client.force_login(self.receptionist)

        response = self.client.get(reverse("day_close"))

        self.assertNotIn("expected_cash", response.context[-1].flatten())
        self.assertNotContains(response, "Expected cash")
        self.assertNotContains(response, "15000")

    def test_day_close_blocks_a_second_close_for_today(self):
        self.client.force_login(self.receptionist)
        self.client.post(reverse("day_close"), {"counted_cash": "0"})

        response = self.client.post(
            reverse("day_close"), {"counted_cash": "0"}, follow=True
        )

        self.assertEqual(DayClose.objects.count(), 1)
        self.assertContains(response, "Today has already been closed")

    def test_day_close_saves_difference(self):
        self.make_payment()
        self.client.force_login(self.receptionist)

        response = self.client.post(
            reverse("day_close"), {"counted_cash": "14000", "note": "Cash counted"}
        )

        self.assertRedirects(response, reverse("day_close"))
        day_close = DayClose.objects.get()
        self.assertEqual(day_close.expected_cash, 15000)
        self.assertEqual(day_close.counted_cash, 14000)
        self.assertEqual(day_close.difference, -1000)
        self.assertTrue(AuditLog.objects.filter(action="DAY_CLOSE").exists())

    def test_owner_report_calculates_known_values(self):
        valid_payment = self.make_payment()
        voided_payment = record_payment(self.visit, [self.consultation], self.owner)
        void_payment(voided_payment, self.owner, "Wrong selection")
        register_reprint(valid_payment, self.receptionist)
        register_reprint(valid_payment, self.receptionist)
        DayClose.objects.create(
            date=timezone.localdate(),
            expected_cash=15000,
            counted_cash=14000,
            difference=-1000,
            closed_by=self.owner,
        )
        self.client.force_login(self.owner)

        response = self.client.get(reverse("reports"))

        self.assertContains(response, "Visits")
        self.assertContains(response, "TZS 15000")
        self.assertContains(response, "TZS 10000")
        self.assertContains(response, "Reprints:</strong> 2")
        self.assertContains(response, "Consultation")
        self.assertContains(response, "Blood test")
        self.assertContains(response, "receptionist")
        self.assertContains(response, "TZS -1000")

    def test_report_detects_missing_receipt_numbers(self):
        Payment.objects.create(
            receipt_no=3,
            visit=self.visit,
            total=0,
            received_by=self.receptionist,
        )
        self.client.force_login(self.owner)

        response = self.client.get(reverse("reports"))

        self.assertContains(response, "Missing: 1, 2")

    def test_report_export_contains_summary_and_sections(self):
        self.make_payment()
        self.client.force_login(self.owner)

        response = self.client.get(reverse("reports_export"))

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(
            workbook.sheetnames,
            ["Summary", "Revenue by day", "Services", "Receptionists", "Receipt gaps", "Day closes"],
        )
        self.assertEqual(workbook["Summary"]["B6"].value, 15000)

    def test_receptionist_gets_403_for_owner_pages(self):
        self.client.force_login(self.receptionist)

        for url_name in ("reports", "reports_export", "audit_log"):
            with self.subTest(url_name=url_name):
                self.assertEqual(self.client.get(reverse(url_name)).status_code, 403)

    def test_successful_and_failed_logins_are_audited(self):
        login_user = User.objects.create_user("login-user", password="safe-password-123")

        self.assertTrue(self.client.login(username="login-user", password="safe-password-123"))
        self.client.logout()
        self.assertFalse(self.client.login(username="login-user", password="incorrect"))

        self.assertTrue(
            AuditLog.objects.filter(user=login_user, action="LOGIN", detail="Successful login").exists()
        )
        self.assertTrue(
            AuditLog.objects.filter(user=None, action="LOGIN", detail__contains="Failed login attempt").exists()
        )

    def test_dashboard_warns_when_local_backup_is_over_24_hours_old(self):
        with TemporaryDirectory() as directory:
            timestamp_file = Path(directory) / "last_backup.txt"
            timestamp_file.write_text((timezone.now() - timedelta(hours=25)).isoformat(), encoding="utf-8")
            self.client.force_login(self.receptionist)
            with patch("reception.views.BACKUP_TIMESTAMP_PATH", timestamp_file):
                response = self.client.get(reverse("dashboard"))

        self.assertContains(response, "Last backup is more than 24 hours old")

    def test_visit_payment_and_receipt_flow(self):
        self.client.force_login(self.receptionist)

        response = self.client.post(
            reverse("visit_new"),
            {"name": "Neema", "phone": "", "age": "28", "sex": "F"},
        )

        visit = Visit.objects.get(name="Neema")
        self.assertRedirects(response, reverse("visit_pay", args=[visit.pk]))
        self.assertEqual(visit.created_by, self.receptionist)
        self.assertTrue(AuditLog.objects.filter(action="VISIT", user=self.receptionist).exists())

        response = self.client.post(
            reverse("visit_pay", args=[visit.pk]),
            {"services": [self.consultation.pk, self.test_service.pk], "price": "1"},
        )

        payment = Payment.objects.get(visit=visit)
        self.assertRedirects(
            response, f"{reverse('receipt_detail', args=[payment.receipt_no])}?print=1"
        )
        self.assertEqual(payment.total, 15000)
        receipt = self.client.get(reverse("receipt_detail", args=[payment.receipt_no]))
        self.assertContains(receipt, "Receipt: 000001")
        self.assertContains(receipt, "Neema")
        self.assertContains(receipt, "TOTAL")
        self.assertContains(receipt, "Cash received")

    def test_visit_form_validates_required_fields(self):
        self.client.force_login(self.receptionist)

        response = self.client.post(reverse("visit_new"), {"name": "", "age": "", "sex": ""})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(Visit.objects.filter(name="").count(), 0)

    def test_visit_pay_and_receipt_require_login(self):
        payment = self.make_payment()

        for url in (
            reverse("visit_new"),
            reverse("visit_pay", args=[self.visit.pk]),
            reverse("receipt_detail", args=[payment.receipt_no]),
        ):
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 302)
                self.assertIn(reverse("login"), response.url)


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class UserManagementTests(TestCase):
    def setUp(self):
        self.owner = User.objects.create_superuser("owner", password="owner-password-123")
        self.receptionist_group = Group.objects.create(name="Receptionist")
        self.receptionist = User.objects.create_user("receptionist", password="reception-password-123")
        self.receptionist.groups.add(self.receptionist_group)

    def test_receptionist_gets_403_for_every_users_url(self):
        self.client.force_login(self.receptionist)
        urls = (
            reverse("user_list"),
            reverse("user_new"),
            reverse("user_password_reset", args=[self.owner.pk]),
            reverse("user_toggle", args=[self.owner.pk]),
        )
        for url in urls:
            with self.subTest(url=url):
                response = self.client.post(url) if url.endswith("/toggle/") else self.client.get(url)
                self.assertEqual(response.status_code, 403)

    def test_owner_creates_receptionist_who_can_log_in(self):
        self.client.force_login(self.owner)
        response = self.client.post(
            reverse("user_new"),
            {
                "username": "new-receptionist",
                "first_name": "Neema",
                "password1": "new-password-123",
                "password2": "new-password-123",
            },
        )

        self.assertRedirects(response, reverse("user_list"))
        created = User.objects.get(username="new-receptionist")
        self.assertTrue(created.is_active)
        self.assertFalse(created.is_staff)
        self.assertFalse(created.is_superuser)
        self.assertTrue(created.groups.filter(name="Receptionist").exists())
        self.assertTrue(self.client.login(username="new-receptionist", password="new-password-123"))
        self.assertTrue(AuditLog.objects.filter(action="USER_CREATE", user=self.owner).exists())

    def test_owner_password_reset_works_and_is_audited(self):
        self.client.force_login(self.owner)
        response = self.client.post(
            reverse("user_password_reset", args=[self.receptionist.pk]),
            {"new_password1": "replacement-password-123", "new_password2": "replacement-password-123"},
        )

        self.assertRedirects(response, reverse("user_list"))
        self.assertTrue(self.client.login(username="receptionist", password="replacement-password-123"))
        self.assertTrue(AuditLog.objects.filter(action="PASSWORD_RESET", user=self.owner).exists())

    def test_deactivated_user_cannot_log_in_and_login_page_explains_why(self):
        self.receptionist.is_active = False
        self.receptionist.save(update_fields=["is_active"])

        self.assertFalse(self.client.login(username="receptionist", password="reception-password-123"))
        response = self.client.post(
            reverse("login"), {"username": "receptionist", "password": "reception-password-123"}
        )
        self.assertContains(response, "This account is deactivated.")

    def test_owner_cannot_deactivate_self_or_last_active_owner(self):
        self.client.force_login(self.owner)

        second_owner = User.objects.create_user("second-owner", password="second-owner-password")
        second_owner.groups.add(Group.objects.create(name="Owner"))

        response = self.client.post(reverse("user_toggle", args=[self.owner.pk]), follow=True)

        self.owner.refresh_from_db()
        self.assertTrue(self.owner.is_active)
        self.assertContains(response, "You cannot deactivate your own account.")

        response = self.client.post(reverse("user_toggle", args=[second_owner.pk]))
        self.assertRedirects(response, reverse("user_list"))
        second_owner.refresh_from_db()
        self.assertFalse(second_owner.is_active)

        response = self.client.post(reverse("user_toggle", args=[self.owner.pk]), follow=True)
        self.assertContains(response, "You cannot deactivate the last active owner.")

    def test_own_password_requires_old_password_and_is_audited(self):
        self.client.force_login(self.receptionist)
        response = self.client.post(
            reverse("account_password"),
            {
                "old_password": "wrong-password",
                "new_password1": "changed-password-123",
                "new_password2": "changed-password-123",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Your old password was entered incorrectly.")

        response = self.client.post(
            reverse("account_password"),
            {
                "old_password": "reception-password-123",
                "new_password1": "changed-password-123",
                "new_password2": "changed-password-123",
            },
        )
        self.assertRedirects(response, reverse("dashboard"))
        self.assertTrue(self.client.login(username="receptionist", password="changed-password-123"))
        self.assertTrue(AuditLog.objects.filter(action="PASSWORD_CHANGE", user=self.receptionist).exists())

    def test_receptionist_cannot_open_admin(self):
        self.client.force_login(self.receptionist)
        response = self.client.get("/admin/")
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("admin:login"), response.url)


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class FirstRunSetupTests(TestCase):
    setup_data = {
        "clinic_name": "Upendo Clinic",
        "owner_full_name": "Asha Owner",
        "owner_username": "asha",
        "owner_password1": "owner-password-123",
        "owner_password2": "owner-password-123",
        "receptionist_username": "neema",
        "receptionist_password": "reception-password-123",
    }

    def test_setup_works_when_there_are_no_users_and_creates_roles(self):
        response = self.client.get(reverse("setup"))
        self.assertEqual(response.status_code, 200)

        response = self.client.post(reverse("setup"), self.setup_data)

        self.assertRedirects(response, reverse("login"))
        owner = User.objects.get(username="asha")
        receptionist = User.objects.get(username="neema")
        self.assertTrue(owner.is_superuser)
        self.assertTrue(owner.is_staff)
        self.assertTrue(owner.groups.filter(name="Owner").exists())
        self.assertFalse(receptionist.is_superuser)
        self.assertFalse(receptionist.is_staff)
        self.assertTrue(receptionist.groups.filter(name="Receptionist").exists())
        self.assertEqual(Setting.objects.get(key="clinic_name").value, "Upendo Clinic")
        self.assertTrue(AuditLog.objects.filter(action="SETUP", user=owner).exists())

    def test_setup_returns_404_after_an_account_exists(self):
        User.objects.create_user("existing", password="existing-password-123")

        response = self.client.get(reverse("setup"))

        self.assertEqual(response.status_code, 404)

    def test_empty_install_redirects_all_normal_pages_to_setup(self):
        response = self.client.get(reverse("dashboard"))

        self.assertRedirects(response, reverse("setup"))

    def test_setup_rejects_weak_or_mismatched_passwords(self):
        weak_data = {**self.setup_data, "owner_password1": "short", "owner_password2": "short"}
        response = self.client.post(reverse("setup"), weak_data)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(User.objects.count(), 0)

        mismatch_data = {**self.setup_data, "owner_password2": "different-password-123"}
        response = self.client.post(reverse("setup"), mismatch_data)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(User.objects.count(), 0)

    def test_setup_clinic_never_creates_users(self):
        call_command("setup_clinic")

        self.assertEqual(User.objects.count(), 0)

    def test_reset_password_command_changes_password_and_writes_audit_row(self):
        user = User.objects.create_user("recover", password="old-password-123")
        with patch(
            "reception.password_recovery.getpass.getpass",
            side_effect=["new-password-123", "new-password-123"],
        ):
            call_command("reset_password", "recover")

        self.assertTrue(self.client.login(username="recover", password="new-password-123"))
        self.assertTrue(AuditLog.objects.filter(action="PASSWORD_RESET", user=user).exists())
